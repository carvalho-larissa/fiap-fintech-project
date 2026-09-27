"""
Gera o documento de entrega da Fase 4 (DOCX) e o exporta para PDF.

Fonte única (nada é redigitado à mão):
  - comandos SQL  : database/comandos/fintech-comandos.sql (validado no Oracle — T13)
  - diagramas     : docs/03-modelagem-de-dados/modelo-logico-v2.png e modelo-fisico-v2.png
  - enunciado     : títulos dos itens vêm do próprio .sql (idênticos ao enunciado)

Saída:
  docs/05-banco-de-dados/entrega/Fintech-Fase4-Comandos-SQL-RM571266.docx
  docs/05-banco-de-dados/entrega/Fintech-Fase4-Comandos-SQL-RM571266.pdf   <- enviado ao portal

Uso: uv run --with python-docx --with pymupdf python docs/05-banco-de-dados/entrega/gerar_documento.py
Requer LibreOffice (soffice) para o PDF.
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[3]
SQL = RAIZ / "database/comandos/fintech-comandos.sql"
IMG_LOG = RAIZ / "docs/03-modelagem-de-dados/modelo-logico-v2.png"
IMG_FIS = RAIZ / "docs/03-modelagem-de-dados/modelo-fisico-v2.png"
RELATORIO = RAIZ / "database/testes/evidencias/relatorio-testes.md"
SAIDA = Path(__file__).parent
NOME = "Fintech-Fase4-Comandos-SQL-RM571266"

COR_MARCA = RGBColor(0x15, 0x45, 0x4E)  # teal do protótipo (frontend)
FONTE_TEXTO = "Calibri"
FONTE_CODIGO = "Consolas"

SECOES = {
    "1": "Cadastro",
    "2": "Alteração",
    "3": "Consultas simples",
    "4": "Consultas ordenadas",
    "5": "Consulta para o dashboard",
}


# ---------------------------------------------------------------------------
# Leitura do .sql
# ---------------------------------------------------------------------------
def ler_comandos() -> list[tuple[str, str, str]]:
    """Retorna [(item, titulo, sql)] na ordem do arquivo. SQL sem comentários, com ';'."""
    texto = SQL.read_text(encoding="utf-8")
    partes = re.split(r"^-- (\d\.\d) (.+)$", texto, flags=re.M)
    itens = []
    for i in range(1, len(partes), 3):
        item, titulo, corpo = partes[i], partes[i + 1].strip(), partes[i + 2]
        linhas = [l.rstrip() for l in corpo.splitlines() if not l.lstrip().startswith("--")]
        sql = "\n".join(linhas).strip()
        sql = sql[: sql.index(";") + 1]
        itens.append((item, titulo, sql))
    assert len(itens) == 15, f"esperava 15 comandos, achei {len(itens)}"
    return itens


# ---------------------------------------------------------------------------
# Helpers de formatação
# ---------------------------------------------------------------------------
def sombrear(par, cor_hex: str):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    pPr.append(shd)


def borda_esquerda(par, cor_hex: str):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in {"w:val": "single", "w:sz": "18", "w:space": "8", "w:color": cor_hex}.items():
        left.set(qn(k), v)
    bdr.append(left)
    pPr.append(bdr)


def manter_com_proximo(par):
    par.paragraph_format.keep_with_next = True


def bloco_sql(doc, sql: str):
    """Um parágrafo por comando, com quebras de linha internas: não se divide entre páginas."""
    p = doc.add_paragraph(style="Codigo")
    p.paragraph_format.keep_together = True
    linhas = sql.split("\n")
    for i, linha in enumerate(linhas):
        run = p.add_run(linha)
        if i < len(linhas) - 1:
            run.add_break(WD_BREAK.LINE)
    sombrear(p, "F3F5F7")
    borda_esquerda(p, "15454E")
    return p


def rodape_com_pagina(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema Fintech · Fase 4 · RM 571266 — página ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            run._r.append(it)


def tabela(doc, cabecalho, linhas, larguras_cm):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, txt in zip(t.rows[0].cells, cabecalho):
        c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "15454E")
        tcPr.append(shd)
    for linha in linhas:
        cells = t.add_row().cells
        for c, txt in zip(cells, linha):
            c.text = ""
            par = c.paragraphs[0]
            # trechos entre `crases` em fonte de código
            for i, pedaco in enumerate(re.split(r"`([^`]+)`", txt)):
                r = par.add_run(pedaco)
                r.font.size = Pt(9)
                if i % 2:
                    r.font.name = FONTE_CODIGO
    for row in t.rows:
        for c, w in zip(row.cells, larguras_cm):
            c.width = Cm(w)
    # LibreOffice ignora a largura das células sem grade explícita e layout fixo
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), larguras_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    # tabela pequena não se divide entre páginas: cada linha "mantém com a próxima"
    # e nenhuma linha quebra no meio
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if i < len(t.rows) - 1:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    return t


def paragrafo_rico(doc, texto: str, style="Normal"):
    """Parágrafo com trechos `código` e **negrito**."""
    p = doc.add_paragraph(style=style)
    for pedaco in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", texto):
        if not pedaco:
            continue
        if pedaco.startswith("`"):
            r = p.add_run(pedaco[1:-1])
            r.font.name = FONTE_CODIGO
            r.font.size = Pt(9.5)
        elif pedaco.startswith("**"):
            r = p.add_run(pedaco[2:-2])
            r.bold = True
        else:
            p.add_run(pedaco)
    return p


# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------
def configurar_estilos(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for nome, tam in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        h = st[nome]
        h.font.name = FONTE_TEXTO
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = COR_MARCA
        h.paragraph_format.space_before = Pt(14 if nome == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        # fonte East Asian/complex também (evita fallback do LibreOffice)
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), FONTE_TEXTO)
    cod = st.add_style("Codigo", 1)
    cod.base_style = st["Normal"]
    cod.font.name = FONTE_CODIGO
    cod.font.size = Pt(8.5)
    cod.paragraph_format.line_spacing = 1.0
    cod.paragraph_format.space_before = Pt(2)
    cod.paragraph_format.space_after = Pt(10)
    cod.paragraph_format.left_indent = Cm(0.2)
    rpr = cod.element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), FONTE_CODIGO)
    rpr.append(rf)


def capa(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FIAP")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = COR_MARCA
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Análise e Desenvolvimento de Sistemas")
    r.font.size = Pt(12)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema Fintech")
    r.bold = True
    r.font.size = Pt(26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fase 4 — Comandos SQL para manipulação de dados (Oracle)")
    r.font.size = Pt(14)
    r.font.color.rgb = COR_MARCA
    for _ in range(8):
        doc.add_paragraph()
    for linha in ("Larissa Gomes de Carvalho", "RM 571266 · Turma 1TDSOA",
                  f"São Paulo, {data_extenso(dt.date.today())}"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run(linha).font.size = Pt(12)
    doc.add_page_break()


def data_extenso(d: dt.date) -> str:
    meses = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]
    return f"{meses[d.month - 1]} de {d.year}"


def resumo_testes() -> tuple[int, int]:
    if not RELATORIO.exists():
        return 0, 0
    m = re.search(r"\*\*(\d+)/(\d+) passaram\*\*", RELATORIO.read_text(encoding="utf-8"))
    return (int(m.group(1)), int(m.group(2))) if m else (0, 0)


def gerar_docx(caminho: Path, comandos):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    configurar_estilos(doc)
    core = doc.core_properties
    core.author = "Larissa Gomes de Carvalho"
    core.title = "Sistema Fintech — Fase 4 — Comandos SQL"
    core.subject = "FIAP 1TDSOA — RM 571266"

    capa(doc)
    rodape_com_pagina(sec)

    # 1. Introdução
    doc.add_heading("1. Introdução", level=1)
    paragrafo_rico(doc, "Este documento reúne os comandos SQL que o sistema **Fintech** — aplicação Java de "
                        "gestão financeira pessoal — utilizará para manipular os dados no banco **Oracle**: "
                        "cadastro, alteração, consultas simples, consultas ordenadas e a consulta do dashboard.")
    paragrafo_rico(doc, "Os comandos foram escritos sobre o modelo relacional criado na Fase 3, evoluído "
                        "conforme permitido no enunciado (\"podendo adaptá-la ou evoluí-la conforme as "
                        "necessidades do seu projeto\"). As evoluções estão listadas na seção 2.3.")
    ok, total = resumo_testes()
    if total:
        paragrafo_rico(doc, f"Todos os comandos foram **executados e validados em um banco Oracle** com uma "
                            f"massa de dados de teste antes da entrega ({ok} de {total} casos de teste "
                            f"aprovados — detalhes no Anexo A).")

    # 2. Modelo relacional
    doc.add_heading("2. Modelo relacional", level=1)
    paragrafo_rico(doc, "Diagramas do modelo relacional do Sistema Fintech (prefixo `T_SF_`). A tabela "
                        "`T_SF_INVESTIMENTO` foi incluída nesta fase para atender às operações de "
                        "investimento solicitadas.")
    for titulo, img in (("2.1 Modelo lógico", IMG_LOG), ("2.2 Modelo físico", IMG_FIS)):
        h = doc.add_heading(titulo, level=2)
        manter_com_proximo(h)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Cm(17))
        legenda = doc.add_paragraph()
        legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = legenda.add_run(f"Figura {titulo[:3]} — {titulo[4:]} (v2, Fase 4). "
                            "Asterisco vermelho = coluna obrigatória.")
        r.italic = True
        r.font.size = Pt(9)

    doc.add_heading("2.3 Evoluções em relação à Fase 3", level=2)
    tabela(doc, ["Evolução", "Motivo"], [
        ["Nova tabela `T_SF_INVESTIMENTO`", "O enunciado exige cadastro, alteração e consulta de investimentos, "
                                            "conceito ausente no modelo da Fase 3."],
        ["Chaves primárias com `IDENTITY`", "O banco gera os códigos; por isso os INSERT não informam a coluna de "
                                            "ID, como nos exemplos do enunciado."],
        ["Domínios com `CHECK`", "Garante valores válidos (ex.: tipo do gasto `FIXO` ou `VARIAVEL`)."],
        ["E-mail único", "Impede dois usuários com o mesmo e-mail (User Story 1)."],
        ["FK composta `(id_conta, id_usuario)`", "Garante que a conta usada em receita, gasto ou investimento "
                                                 "pertence ao mesmo usuário."],
        ["`dt_pagamento` da parcela opcional", "Permite registrar parcelas em aberto ou em atraso (User Story 4)."],
        ["Taxas como fração decimal", "`0.0299` = 2,99%; mantém o tipo `NUMBER(5,4)` da Fase 3 comportando "
                                      "juros de cartão acima de 10% a.m."],
    ], [6, 11])
    paragrafo_rico(doc, "Nenhuma coluna da Fase 3 foi removida ou renomeada.")

    # 3. Convenção
    doc.add_heading("3. Convenção das máscaras de substituição", level=1)
    manter_com_proximo(paragrafo_rico(doc, "Os valores dinâmicos, que a aplicação Java preencherá, "
                                           "aparecem como máscaras:"))
    tabela(doc, ["Tipo do valor", "Máscara", "Exemplo"], [
        ["Texto", "`'[NOME DO CAMPO]'`", "`'[DESCRIÇÃO DO GASTO]'`"],
        ["Numérico", "`[NOME DO CAMPO]`", "`[VALOR DO GASTO]`"],
        ["Data", "`TO_DATE('[DATA]', 'DD/MM/YYYY')`", "`TO_DATE('[DATA DO GASTO]', 'DD/MM/YYYY')`"],
    ], [3, 6.5, 7.5])
    paragrafo_rico(doc, "Valores de domínio são informados em maiúsculas e sem acento (ex.: `'VARIAVEL'`). "
                        "Os comandos não contêm `COMMIT`: o controle da transação fica com a aplicação.")

    # 4. Comandos
    doc.add_heading("4. Comandos SQL", level=1)
    secao_atual = None
    for item, titulo, sql in comandos:
        sec_id = item.split(".")[0]
        if sec_id != secao_atual:
            secao_atual = sec_id
            doc.add_heading(f"4.{sec_id} {SECOES[sec_id]}", level=2)
        h = doc.add_heading(f"4.{item} {titulo[0].upper()}{titulo[1:]}", level=3)
        manter_com_proximo(h)
        bloco_sql(doc, sql)

    # Anexo
    ok, total = resumo_testes()
    if total:
        doc.add_page_break()
        doc.add_heading("Anexo A — Validação dos comandos", level=1)
        paragrafo_rico(doc, "Os comandos da seção 4 foram executados em Oracle Database Free (container Docker), "
                            "com sintaxe mantida compatível com o Oracle 19c. Para cada caso de teste, o comando "
                            "entregue foi usado exatamente como está, apenas com as máscaras substituídas por "
                            "valores de uma massa de dados de teste.")
        tabela(doc, ["Grupo", "Casos", "O que foi verificado"], [
            ["Estrutura", "6", "10 tabelas, objetos válidos, constraints nomeadas, índices, IDENTITY"],
            ["Cadastro", "9", "5 INSERT válidos; rejeição de e-mail duplicado, domínio inválido, "
                              "usuário inexistente e conta de outro usuário"],
            ["Alteração", "7", "cada UPDATE afeta só o registro do usuário informado; senha preservada"],
            ["Consultas simples", "5", "registro correto; 0 linhas para outro usuário; senha nunca retornada"],
            ["Consultas ordenadas", "6", "ordem da mais recente à mais antiga, com desempate por código"],
            ["Dashboard", "7", "sempre 1 linha por usuário, inclusive sem despesas ou investimentos"],
            ["Constraints e compatibilidade", "10", "regras de integridade e ausência de recursos exclusivos "
                                                    "de versões recentes do Oracle"],
        ], [4, 1.5, 11.5])
        p = paragrafo_rico(doc, f"**Resultado: {ok} de {total} casos aprovados.**")
        p.paragraph_format.space_before = Pt(8)

    doc.save(caminho)


# ---------------------------------------------------------------------------
# PDF + verificações
# ---------------------------------------------------------------------------
def soffice() -> str:
    for c in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if c and Path(c).exists():
            return c
    sys.exit("LibreOffice (soffice) não encontrado — instale: winget install TheDocumentFoundation.LibreOffice")


def exportar_pdf(docx: Path) -> Path:
    subprocess.run([soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, capture_output=True, timeout=180)
    pdf = docx.with_suffix(".pdf")
    assert pdf.exists(), "PDF não foi gerado"
    return pdf


def verificar_pdf(pdf: Path, comandos) -> list[str]:
    """Confere que cada comando do .sql aparece no PDF (texto extraído, espaços normalizados)."""
    import pymupdf

    doc = pymupdf.open(pdf)
    texto = "\n".join(p.get_text() for p in doc)
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    corpo = norm(texto)
    problemas = []
    for item, titulo, sql in comandos:
        if norm(sql) not in corpo:
            problemas.append(f"comando {item} não encontrado literalmente no PDF")
        if norm(f"4.{item} {titulo[0].upper()}{titulo[1:]}") not in corpo:
            problemas.append(f"título {item} não encontrado no PDF")
    # aspas tipográficas (autocorreção do Word) tornariam o SQL inválido
    for *_, sql in comandos:
        assert not re.search("[\u2018\u2019\u201c\u201d]", sql), "aspas tipográficas no .sql"
    if doc.page_count < 5:
        problemas.append(f"PDF com poucas páginas ({doc.page_count})")
    return problemas, doc.page_count


def main():
    comandos = ler_comandos()
    for img in (IMG_LOG, IMG_FIS):
        assert img.exists(), f"imagem ausente: {img}"
    docx = SAIDA / f"{NOME}.docx"
    gerar_docx(docx, comandos)
    pdf = exportar_pdf(docx)
    problemas, paginas = verificar_pdf(pdf, comandos)
    print(f"DOCX: {docx}\nPDF : {pdf} ({paginas} páginas)")
    if problemas:
        print("PROBLEMAS:\n  " + "\n  ".join(problemas))
        sys.exit(1)
    print("Verificação: os 15 comandos e títulos aparecem literalmente no PDF (aspas retas preservadas).")


if __name__ == "__main__":
    main()
